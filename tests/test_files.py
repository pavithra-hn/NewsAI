import io

import docx
import pytest

from demo import files


def word_file(build) -> bytes:
    document = docx.Document()
    build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def pdf_file(text: str) -> bytes:
    """A one-page PDF holding one line of text, built by hand."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b"BT ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def test_a_text_file_is_read_as_utf8():
    data = "Engine: 286 hp\nمحرك قوي".encode()

    assert files.read_upload("notes.txt", data) == "Engine: 286 hp\nمحرك قوي"


def test_a_text_file_loses_its_byte_order_mark():
    assert files.read_upload("notes.txt", "﻿Hello".encode()) == "Hello"


def test_a_word_file_keeps_its_paragraphs_lists_and_tables_in_order():
    def build(document):
        document.add_paragraph("1. Common Specifications")
        document.add_paragraph("Engine: 286 hp", style="List Bullet")
        document.add_paragraph("")
        document.add_paragraph("Brakes: check the wear indicators", style="List Number")
        document.add_paragraph("Fluids: review the oil reports", style="List Number")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Fuel tank"
        table.cell(0, 1).text = "380 L"
        document.add_paragraph("Last line")

    text = files.read_upload("spec.docx", word_file(build))

    assert text.splitlines() == [
        "1. Common Specifications",
        "• Engine: 286 hp",
        "1. Brakes: check the wear indicators",
        "2. Fluids: review the oil reports",
        "Fuel tank | 380 L",
        "Last line",
    ]


def test_a_numbered_list_starts_again_after_ordinary_text():
    def build(document):
        document.add_paragraph("First", style="List Number")
        document.add_paragraph("Between")
        document.add_paragraph("Second list", style="List Number")

    lines = files.read_upload("list.docx", word_file(build)).splitlines()

    assert lines == ["1. First", "Between", "1. Second list"]


def test_a_pdf_gives_its_text():
    text = files.read_upload("spec.pdf", pdf_file("Operating weight 23125 kg"))

    assert "Operating weight 23125 kg" in text


def test_a_pdf_without_text_is_refused_as_a_probable_scan():
    with pytest.raises(files.ExtractError, match="scan"):
        files.read_upload("scan.pdf", pdf_file(""))


def test_an_unsupported_file_type_names_the_supported_ones():
    with pytest.raises(files.ExtractError, match=r"\.docx.*\.pdf.*\.txt"):
        files.read_upload("photo.png", b"not text")


@pytest.mark.parametrize("name", ["broken.docx", "broken.pdf"])
def test_a_damaged_file_is_refused_with_a_plain_message(name):
    with pytest.raises(files.ExtractError, match="could not be read"):
        files.read_upload(name, b"this is not a real document")


def test_the_file_type_is_read_whatever_the_case_of_its_name():
    assert files.read_upload("NOTES.TXT", b"Hello") == "Hello"


def test_a_word_file_made_from_text_reads_back_the_same():
    text = "Section one\n• Engine: 286 hp\n• Weight: 23,125 kg\n1. Check the brakes\n2. Check the tires"

    assert files.read_upload("out.docx", files.to_docx(text)) == text


def test_a_word_file_keeps_arabic_text():
    text = "المواصفات\n• المحرك: 286 حصاناً"

    assert files.read_upload("out.docx", files.to_docx(text)) == text


# Found writing the 966H download: Word numbers every "List Number" paragraph as
# one list, so section headings 1 to 5 pushed the checklist on to 6 to 13.


def word_paragraphs(data: bytes):
    return [(p.style.name, p.text) for p in docx.Document(io.BytesIO(data)).paragraphs]


def test_a_word_download_keeps_every_number_exactly_as_written():
    text = "1. Section one\n\u2022 Item\n2. Section two\n1. Check the brakes\n2. Check the tires"

    paragraphs = word_paragraphs(files.to_docx(text))

    assert [style for style, _ in paragraphs if style.startswith("List Number")] == []
    assert [t for _, t in paragraphs] == [
        "1. Section one", "Item", "2. Section two", "1. Check the brakes", "2. Check the tires",
    ]


def test_a_bullet_followed_by_a_tab_becomes_a_word_bullet():
    paragraphs = word_paragraphs(files.to_docx("\u2022\tEngine: 286 hp"))

    assert paragraphs == [("List Bullet", "Engine: 286 hp")]


def test_a_text_with_tabbed_bullets_reads_back_with_the_same_content():
    text = "Intro line\n\u2022\tEngine: 286 hp\n1.\tBrakes: check wear"

    back = files.read_upload("out.docx", files.to_docx(text))

    assert back.splitlines() == ["Intro line", "\u2022 Engine: 286 hp", "1.\tBrakes: check wear"]
