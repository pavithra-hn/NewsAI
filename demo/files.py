"""Reading an uploaded document as plain text, and writing text back as Word.

Lists are kept as lines that start with "• " or "1. ", so the humanizer sees the
same structure the reader does and the rewrite can keep it.
"""

import io
import re
import zipfile
from pathlib import PurePath

import docx
from docx.document import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED = (".docx", ".pdf", ".txt")
BULLET = "• "
NUMBERED = re.compile(r"^(\d+)\.\s+(.*)$")


class ExtractError(ValueError):
    """The upload could not be turned into text. The message is shown as it is."""


def read_upload(name: str, data: bytes) -> str:
    suffix = PurePath(name).suffix.lower()
    if suffix == ".txt":
        return data.decode("utf-8-sig", errors="replace").replace("\r\n", "\n").strip()
    if suffix == ".docx":
        return _read_docx(data)
    if suffix == ".pdf":
        return _read_pdf(data)
    raise ExtractError(f"This file type is not supported. Upload a {', '.join(SUPPORTED)} file.")


def _read_docx(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractError("This Word file could not be read.") from exc

    lines, number = [], 0
    for block in _blocks(document):
        if isinstance(block, Table):
            number = 0
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        if style.startswith("List Number"):
            number += 1
            lines.append(f"{number}. {text}")
            continue
        number = 0
        lines.append(BULLET + text if style.startswith("List") else text)
    return "\n".join(lines)


def _blocks(document: Document):
    """Paragraphs and tables in the order they appear in the document."""
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, document)
        elif tag == "tbl":
            yield Table(child, document)


def _read_pdf(data: bytes) -> str:
    try:
        pages = [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise ExtractError("This PDF could not be read.") from exc
    text = "\n\n".join(page.strip() for page in pages if page.strip())
    if not text:
        raise ExtractError(
            "No text was found in this PDF. It may be a scan; paste the text instead."
        )
    return text


def to_docx(text: str) -> bytes:
    """A Word file of the text, with "• " and "1. " lines as real lists."""
    document = docx.Document()
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        numbered = NUMBERED.match(line)
        if line.startswith(BULLET):
            document.add_paragraph(line[len(BULLET):], style="List Bullet")
        elif numbered:
            document.add_paragraph(numbered.group(2), style="List Number")
        else:
            document.add_paragraph(line)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
